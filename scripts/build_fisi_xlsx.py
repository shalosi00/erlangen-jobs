#!/usr/bin/env python3
"""Build styled XLSX deliverable from the merged FISI postings CSV.
Sheets: Postings (sorted by date desc) / Summary (live COUNTIF) / Review (cross-checks)."""
import sys, os, csv, datetime

XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path: sys.path.insert(0, sub)
import base  # design tokens - single source of truth
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

CSV_IN = "/home/z/my-project/download/fisi_junior_erlangen50km_last14d.csv"
OUT = "/home/z/my-project/download/FISI_Junior_Postings_Erlangen50km_2weeks.xlsx"

rows = list(csv.DictReader(open(CSV_IN, encoding="utf-8")))
rows.sort(key=lambda r: r["date_posted"], reverse=True)
n = len(rows)

wb = Workbook()
wb.properties.creator = "Z.ai"

# ---------------- Sheet 1: Postings ----------------
ws = wb.active
ws.title = "Postings"
ws["A1"] = "FISI / Junior IT - First Job after Ausbildung"
ws["A1"].font = Font(name=base.FONT_NAME, size=14, bold=True, color=base.NEUTRAL_900)
ws["A2"] = ("Erlangen + 50 km - posted within last 14 days - "
            f"generated {datetime.date.today().isoformat()} - sources: BA Jobsuche v6, Indeed DE (JobSpy), LinkedIn guest, Arbeitnow")
ws["A2"].font = Font(name=base.FONT_NAME, size=9, italic=True, color=base.NEUTRAL_600)

headers = ["Date", "Title", "Company", "City", "PLZ", "Source", "Salary Min",
           "Salary Max", "Currency", "Remote", "Junior", "URL", "Ref Nr"]
HR = 4  # header row
for j, h in enumerate(headers, 2):
    ws.cell(row=HR, column=j, value=h)
base.style_header_row(ws, HR, 2, 1 + len(headers))

for i, r in enumerate(rows):
    row = HR + 1 + i
    vals = [datetime.date.fromisoformat(r["date_posted"]) if r["date_posted"] else None,
            r["title"], r["company"], r["city"], r["plz"], r["source"],
            (float(r["salary_min"]) if r["salary_min"] else None),
            (float(r["salary_max"]) if r["salary_max"] else None),
            r["currency"], r["remote"], r["junior_signal"], r["url"], r["refnr"]]
    for j, v in enumerate(vals, 2):
        ws.cell(row=row, column=j, value=v)
    base.style_header_row  # noqa (keep linters quiet)
    # alternating row fill
    try:
        fill = base.fill_data_row(i)
        if fill:
            for j in range(2, 2 + len(headers)):
                ws.cell(row=row, column=j).fill = fill
    except Exception:
        pass
    d = ws.cell(row=row, column=2)
    if r["date_posted"]:
        d.number_format = "YYYY-MM-DD"
    for c_num in (8, 9):
        ws.cell(row=row, column=c_num).number_format = "#,##0"
        ws.cell(row=row, column=c_num).alignment = base.align_number()
    ws.cell(row=row, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    if r["url"]:
        c = ws.cell(row=row, column=13)
        c.hyperlink = r["url"]
        c.value = "open"
        c.font = Font(name=base.FONT_NAME, size=10, color="0563C1", underline="single")
    if r["junior_signal"] == "yes":
        ws.cell(row=row, column=12).fill = base.CF_POSITIVE_FILL
        ws.cell(row=row, column=12).font = base.CF_POSITIVE_FONT

END = HR + n
ws.freeze_panes = f"B{HR + 1}"
ws.auto_filter.ref = f"B{HR}:{get_column_letter(1 + len(headers))}{END}"
base.auto_fit_columns(ws, min_width=9, max_width=52, header_row=HR, data_start_row=HR + 1)
ws.column_dimensions["C"].width = 52

# ---------------- Sheet 2: Summary ----------------
sm = wb.create_sheet("Summary")
sm["A1"] = "Summary - live formulas (update if you edit the Postings sheet)"
sm["A1"].font = Font(name=base.FONT_NAME, size=12, bold=True, color=base.NEUTRAL_900)

def block(ws2, r0, title, label_col_values, countif_col, wildcard=False):
    ws2.cell(row=r0, column=1, value=title).font = Font(
        name=base.FONT_NAME, size=11, bold=True, color=base.PRIMARY)
    hdr = r0 + 1
    ws2.cell(row=hdr, column=1, value="Value")
    ws2.cell(row=hdr, column=2, value="Count")
    base.style_header_row(ws2, hdr, 1, 2)
    for k, val in enumerate(label_col_values):
        rr = hdr + 1 + k
        ws2.cell(row=rr, column=1, value=val)
        crit = f"*{val}*" if wildcard else val
        ws2.cell(row=rr, column=2,
                 value=f'=COUNTIF(Postings!${countif_col}${HR + 1}:${countif_col}${END},"{crit}")')
        ws2.cell(row=rr, column=2).alignment = base.align_number()
    tot = hdr + 1 + len(label_col_values)
    ws2.cell(row=tot, column=1, value="TOTAL").font = Font(name=base.FONT_NAME, bold=True)
    ws2.cell(row=tot, column=2,
             value=f"=SUM(B{hdr + 1}:B{tot - 1})").font = Font(name=base.FONT_NAME, bold=True)
    ws2.cell(row=tot, column=2).alignment = base.align_number()
    return tot + 2

sources = ["BA", "Indeed", "LinkedIn"]  # wildcard: combos counted in each
cities = sorted({r["city"] for r in rows if r["city"]})
dates = sorted({r["date_posted"] for r in rows if r["date_posted"]}, reverse=True)
r0 = block(sm, 3, "By source (wildcard: multi-source rows count in each)", sources, "G", wildcard=True)
r1 = block(sm, r0, "By city", cities, "E")
r2 = block(sm, r1, "By date posted", dates, "B")
sm.cell(row=r2, column=1, value="Junior signal rows").font = Font(name=base.FONT_NAME, bold=True)
sm.cell(row=r2, column=2, value=f'=COUNTIF(Postings!$L${HR + 1}:$L${END},"yes")').alignment = base.align_number()
base.auto_fit_columns(sm, min_width=10, max_width=34, header_row=4)

# ---------------- Sheet 3: Review ----------------
rv = wb.create_sheet("Review")
rv.sheet_properties.tabColor = "FFC000"
rv["A1"] = "Cross-validation checks"
rv["A1"].font = Font(name=base.FONT_NAME, size=12, bold=True)
checks = [
    ["Check", "Expected", "Actual", "Status"],
    ["Posting count", n, f"=COUNTA(Postings!C{HR + 1}:C{END})",
     '=IF(B2=C2,"PASS","FAIL")'],
    ["Junior-signal rows", sum(1 for r in rows if r["junior_signal"] == "yes"),
     f'=COUNTIF(Postings!L{HR + 1}:L{END},"yes")', '=IF(B3=C3,"PASS","FAIL")'],
    ["Oldest posting within 14 days", "TRUE",
     f"=MIN(Postings!B{HR + 1}:B{END})>=DATE(YEAR(TODAY()),MONTH(TODAY()-14),DAY(TODAY()-14))",
     '=IF(C4=TRUE,"PASS","FAIL")'],
]
for i, row in enumerate(checks, 1):
    for j, val in enumerate(row, 1):
        rv.cell(row=i, column=j, value=val)
base.style_header_row(rv, 1, 1, 4)
base.auto_fit_columns(rv, min_width=10, max_width=44, header_row=1)

wb.save(OUT)
print(f"saved {OUT} | postings: {n} | cities: {len(cities)} | dates: {len(dates)}")
